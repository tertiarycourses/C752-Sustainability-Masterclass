"""Post-conversion fix-ups for the C752 deck: strip the WSQ layer convert_deck.py left behind."""
import re
from pptx import Presentation

PATH = "courseware/Sustainability Masterclass (C752)-v1.0.pptx"
p = Presentation(PATH)


def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise SystemExit(f"shape {sid} not found")


def set_para(para, line):
    runs = [r for r in para.runs if r.text] or para.runs
    runs[0].text = line
    for r in para.runs:
        if r._r is not runs[0]._r:
            r.text = ""


def set_paras(slide, sid, lines):
    paras = shape(slide, sid).text_frame.paragraphs
    assert len(paras) == len(lines), (sid, len(paras), lines)
    for para, line in zip(paras, lines):
        set_para(para, line)


def find(text):
    hits = [s for s in p.slides if any(sh.has_text_frame and sh.text_frame.text == text for sh in s.shapes)]
    assert len(hits) == 1, (text, len(hits))
    return hits[0]


# ---- cover
cover = p.slides[0]
set_paras(cover, 9, ["COURSE SLIDES  ·  NON-WSQ"])
code = shape(cover, 12).text_frame.paragraphs
set_para(code[0], "Course Code: C752")

# ---- administration divider
set_paras(p.slides[1], 6, ["Welcome and Outcomes"])

# ---- trainer card role line
for s in p.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                if "".join(r.text for r in para.runs) == "Course facilitation, assessment briefing and learner support":
                    set_para(para, "Course facilitation, activity coaching and learner support")

# ---- learning resources: the open-book assessment card
res = find("Access Your Course Materials")
set_paras(res, 36, ["Reference", "Keep the Learner Guide and source register to hand"])

# ---- course map + topic/content eyebrows: drop the K#/A# competency codes
KA = re.compile(r"(?:\s*·\s*[KA]\d+)+$")      # "Topic 1 ... · K1 · K3 · K7 · A1", "TOPIC 1 · K1 · K3 · K7 · A1"
LEAD = re.compile(r"^[KA]\d+\s*·\s*")          # "K1 · STANDARDS"
for s in p.slides:
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs)
            new = LEAD.sub("", KA.sub("", t))
            if new != t:
                set_para(para, new)

# ---- recommended courses: course names without the WSQ prefix
rec = find("Recommended Courses")
for sh in rec.shapes:
    if sh.has_text_frame:
        for para in sh.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs)
            if t.startswith("WSQ "):
                set_para(para, t[4:])

# ---- one 30-minute lunch (no separate tea breaks in a 9:30am-5:30pm non-WSQ day)
lunch = find("LUNCH BREAK")
for sh in lunch.shapes:
    if sh.has_text_frame and sh.text_frame.text == "60 minutes":
        set_para(sh.text_frame.paragraphs[0], "30 minutes")

# ---- Course Feedback: the digital-attendance slide rewritten in place
fb = p.slides[-2]
set_paras(fb, 4, ["COURSE FEEDBACK"])
set_paras(fb, 5, ["Course Feedback & Certificate"])
set_paras(fb, 11, ["Feedback", "Complete the course feedback survey before you leave."])
set_paras(fb, 16, ["Improve", "Tell us what worked and what to improve."])
set_paras(fb, 21, ["Support", "Ask the trainer if you cannot open the survey or files."])
set_paras(fb, 26, ["Certificate", "Your certificate of completion will be emailed to you."])

# ---- closing slide eyebrow carried the old course name
set_paras(p.slides[-1], 4, ["SUSTAINABILITY MASTERCLASS"])

p.core_properties.title = "Sustainability Masterclass"
p.core_properties.subject = "C752"
p.save(PATH)
print("saved")
